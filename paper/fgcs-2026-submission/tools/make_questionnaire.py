#!/usr/bin/env python3
"""
Build the Word file answering the FGCS Editorial Manager questionnaire
(items 1 to 3), each answer within the 1500-character limit.

Usage:  python tools/make_questionnaire.py
"""

import os
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.abspath(os.path.join(HERE, ".."))
FONT = "Times New Roman"
LIMIT = 1500

TITLE = ("CR-SHE: Oblivious homomorphic aggregation for collusion-resilient "
         "search and analytics over encrypted multi-cloud IoT data")
JOURNAL = "Future Generation Computer Systems"

Q1 = """Research domains: applied cryptography; privacy-preserving distributed and multi-cloud systems; encrypted search and data management for the Internet of Things.

The manuscript addresses oblivious keyword search and analytics over IoT telemetry replicated across several independent cloud providers. Its contributions are:

(1) Oblivious homomorphic aggregation. Each provider folds its own distributed point function (DPF) share directly into a homomorphically encrypted record column and returns a single ciphertext, so the matched record set is never assembled at any provider and its cardinality never leaves the client.

(2) A precomputed-column fast path that makes single-keyword aggregates over templates registered at setup independent of the record count.

(3) A one-time-pad index mask, cancelled inside the ciphertext, that removes posting-list-length leakage from the replicated index at rest.

(4) Proofs of correctness and of adaptive leakage security against a coalition of up to t = n-1 colluding providers, with a per-query coalition leakage profile revealing only the requested aggregate class, plus an explicit algebraic compatibility condition for composing DPF shares with the homomorphic plaintext ring.

(5) A C++ prototype evaluated on two real IoT corpora against ORAM, single-server searchable encryption, distributed-trust search and a faithful homomorphic equality scan, with an empirical leakage-abuse attack."""

Q2 = """This manuscript is not an extension of previously published work by the authors. No conference or workshop version exists, and it has not been posted as a preprint.

Distributed-trust query systems already hide which records a query selects. The novelty here is the composition step they leave open. Selection and computation are not the same problem: once a scheme aggregates over the matched records, it typically hands the matched set to the providers, and a colluding coalition then recovers exactly the access pattern the selection phase protected.

The novel aspects are:

(1) Oblivious homomorphic aggregation, in which the analytics phase inherits the access-pattern hiding of the selection phase, over data encrypted at rest under a single owner's key, with no server-to-server channel and a collusion bound of t = n-1. By contrast, Splinter computes over plaintext databases, DORY stops at retrieval, and Waldo assumes an honest majority (t = 1) with secret-shared data and inter-server interaction.

(2) An index mask compatible with oblivious aggregation: the mask term is cancelled inside the ciphertext via an encrypted correction column retrieved by the same DPF, so no provider ever sees an unmasked row or a mask value.

(3) An algebraic compatibility condition made explicit, whose violation is silent and data-dependent.

(4) A leakage-abuse attack measured against our own index rather than an assertion that the residual leakage is benign."""

Q3 = """Future Generation Computer Systems covers distributed systems, cloud infrastructures and the Internet of Things, and names security aspects and information management among its methods and tools. This manuscript sits in that intersection.

It treats multi-cloud replication, which deployments already adopt for availability and vendor independence, as a systems property to be exploited for privacy rather than merely tolerated, and asks what a coalition of providers can infer when the same encrypted data sits at all of them.

The work is systems-oriented rather than purely theoretical. Alongside the security proofs, it reports an end-to-end C++ prototype running n independent provider processes and a client over TCP, measured on two public IoT corpora of 405,184 and 420,370 records. A registered single-keyword aggregate costs 39 ms per provider irrespective of record count, provider compute is flat in n, per-query downlink is fixed at 526 kB independently of the result-set size, and throughput plateaus near 125 queries per second. The paper also identifies where the design becomes bounded by storage rather than time, and states plainly what the evaluation does not establish, including wide-area deployment and gateway-class hardware.

The IoT data-management setting, the multi-cloud infrastructure model, the encrypted index and analytics, and the measured performance study all fall within the journal's stated scope."""

QUESTIONS = [
    ("1. Briefly list the research domain(s) and the contributions of the "
     "manuscript", Q1),
    ("2. Briefly describe the novel aspects of the manuscript. Authors whose "
     "manuscript extends work previously published should identify the parts "
     "of their submission that contain novel contributions", Q2),
    ("3. Briefly explain how the manuscript matches the journal's scope", Q3),
]


def para(doc, text="", size=11, bold=False, italic=False, align=None,
         before=0, after=8, indent=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    if align is not None:
        p.alignment = align
    if indent is not None:
        p.paragraph_format.left_indent = Inches(indent)
    if text:
        r = p.add_run(text)
        r.bold = bold
        r.italic = italic
        r.font.size = Pt(size)
        r.font.name = FONT
    return p


def rule(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(10)
    b = OxmlElement("w:pBdr")
    bot = OxmlElement("w:bottom")
    bot.set(qn("w:val"), "single")
    bot.set(qn("w:sz"), "6")
    bot.set(qn("w:space"), "1")
    bot.set(qn("w:color"), "999999")
    b.append(bot)
    p._p.get_or_add_pPr().append(b)


def build():
    over = []
    for head, ans in QUESTIONS:
        n = len(ans)
        print("%-4s %5d / %d chars %s" % (head[:2], n, LIMIT,
                                          "OK" if n <= LIMIT else "OVER"))
        if n > LIMIT:
            over.append(head[:2])
    if over:
        print("ERROR: over the limit:", over)
        return None

    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = FONT
    st.font.size = Pt(11)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    st.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    st.paragraph_format.space_after = Pt(0)
    for s in doc.sections:
        s.top_margin = s.bottom_margin = Inches(1.0)
        s.left_margin = s.right_margin = Inches(1.1)

    para(doc, "Submission questionnaire", size=14, bold=True, after=4)
    para(doc, "Submitted to %s" % JOURNAL, size=10, italic=True, after=2)
    para(doc, TITLE, size=10, italic=True, after=4)
    rule(doc)
    para(doc, "Paste each answer into the matching box in Editorial Manager. "
              "Character counts are given for reference; the limit is 1500 "
              "per answer.", size=10, italic=True, after=14)

    for head, ans in QUESTIONS:
        para(doc, head, size=11.5, bold=True, after=6)
        for block in ans.split("\n\n"):
            para(doc, block, size=11, align=WD_ALIGN_PARAGRAPH.JUSTIFY,
                 after=6, indent=0.15)
        para(doc, "[%d of %d characters]" % (len(ans), LIMIT), size=9,
             italic=True, after=16, indent=0.15)

    out = os.path.join(ROOT, "CRSHE_FGCS_SubmissionQuestionnaire.docx")
    doc.save(out)
    return out


if __name__ == "__main__":
    o = build()
    if o:
        print("wrote", o)
    else:
        sys.exit(1)
