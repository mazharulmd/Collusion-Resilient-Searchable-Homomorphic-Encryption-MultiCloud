#!/usr/bin/env python3
"""
Build the Word version of the CR-SHE manuscript from CRSHE_FGCS.tex.

There is no pandoc or LibreOffice on this machine, so this is a small
purpose-built converter rather than a general one: it knows the structure of
this particular manuscript. It produces a single-column, double-spaced Word
file with continuous line numbers, which is the shape Elsevier wants for a
Word submission.

Maths is rendered with real Word subscript and superscript runs and Unicode
operators, not as images. Display equations become centred paragraphs.

Usage:  python tools/tex_to_docx.py
"""

import os
import re
import sys

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.abspath(os.path.join(HERE, ".."))
TEX = os.path.join(ROOT, "CRSHE_FGCS.tex")
BIB = os.path.join(ROOT, "references.bib")
OUT = os.path.join(ROOT, "CRSHE_FGCS_Manuscript.docx")

BODY_FONT = "Times New Roman"
MONO_FONT = "Consolas"

# ---------------------------------------------------------------------------
# LaTeX macro expansion
# ---------------------------------------------------------------------------
USER_MACROS = {
    r"\negl": r"\mathsf{negl}", r"\PPT": r"\mathrm{PPT}",
    r"\Real": r"\mathbf{Real}", r"\Ideal": r"\mathbf{Ideal}",
    r"\Leak": r"\mathcal{L}", r"\Sim": r"\mathcal{S}",
    r"\Adv": r"\mathcal{A}", r"\cind": r"\approxc",
    r"\Gen": r"\mathsf{Gen}", r"\Eval": r"\mathsf{Eval}",
    r"\Setupp": r"\mathsf{Setup}", r"\BuildIndex": r"\mathsf{BuildIndex}",
    r"\EncData": r"\mathsf{EncData}", r"\Token": r"\mathsf{Token}",
    r"\Search": r"\mathsf{Search}", r"\Compute": r"\mathsf{Compute}",
    r"\Dec": r"\mathsf{Dec}", r"\Enc": r"\mathsf{Enc}",
    r"\advPRF": r"\mathbf{Adv}^{\mathrm{prf}}",
    r"\advDPF": r"\mathbf{Adv}^{\mathrm{dpf}}",
    r"\advCPA": r"\mathbf{Adv}^{\mathrm{cpa}}",
    r"\bplus": r"\boxplus", r"\bminus": r"\boxminus", r"\bdot": r"\boxdot",
}

SYMBOLS = {
    "alpha": "\u03b1", "beta": "\u03b2", "gamma": "\u03b3", "delta": "\u03b4",
    "epsilon": "\u03b5", "zeta": "\u03b6", "eta": "\u03b7", "theta": "\u03b8",
    "iota": "\u03b9", "kappa": "\u03ba", "lambda": "\u03bb", "mu": "\u03bc",
    "nu": "\u03bd", "xi": "\u03be", "pi": "\u03c0", "rho": "\u03c1",
    "sigma": "\u03c3", "tau": "\u03c4", "phi": "\u03c6", "chi": "\u03c7",
    "psi": "\u03c8", "omega": "\u03c9",
    "Gamma": "\u0393", "Delta": "\u0394", "Theta": "\u0398",
    "Lambda": "\u039b", "Sigma": "\u03a3", "Omega": "\u03a9",
    "Pi": "\u03a0", "Phi": "\u03a6", "Psi": "\u03a8",
    "sum": "\u2211", "prod": "\u220f", "int": "\u222b",
    "in": "\u2208", "notin": "\u2209", "subseteq": "\u2286",
    "subset": "\u2282", "cup": "\u222a", "cap": "\u2229",
    "emptyset": "\u2205", "le": "\u2264", "ge": "\u2265", "leq": "\u2264",
    "geq": "\u2265", "neq": "\u2260", "approx": "\u2248",
    "approxc": "\u2248", "equiv": "\u2261", "sim": "\u223c",
    "times": "\u00d7", "cdot": "\u00b7", "pm": "\u00b1", "mp": "\u2213",
    "rightarrow": "\u2192", "leftarrow": "\u2190", "to": "\u2192",
    "gets": "\u2190", "mapsto": "\u21a6", "Rightarrow": "\u21d2",
    "boxplus": "\u229e", "boxminus": "\u229f", "boxdot": "\u22a1",
    "oplus": "\u2295", "otimes": "\u2297",
    "lceil": "\u2308", "rceil": "\u2309", "lfloor": "\u230a",
    "rfloor": "\u230b", "langle": "\u27e8", "rangle": "\u27e9",
    "dots": "\u2026", "ldots": "\u2026", "cdots": "\u22ef",
    "vdots": "\u22ee", "infty": "\u221e", "partial": "\u2202",
    "forall": "\u2200", "exists": "\u2203", "neg": "\u00ac",
    "checkmark": "\u2713", "dagger": "\u2020", "circ": "\u2218",
    "star": "\u22c6", "surd": "\u221a", "ell": "\u2113",
    "quad": "\u2003", "qquad": "\u2003\u2003",
    "bmod": " mod", "pmod": " mod", "colon": ":",
    "vert": "|", "mid": "|", "backslash": "\\",
}

BLACKBOARD = {"Z": "\u2124", "R": "\u211d", "N": "\u2115", "Q": "\u211a",
              "C": "\u2102", "G": "\U0001d53e", "F": "\U0001d53d"}
CALLIGRAPHIC = {"L": "\u2112", "A": "\U0001d49c", "S": "\U0001d4ae",
                "B": "\u212c", "D": "\U0001d49f", "P": "\U0001d4ab",
                "R": "\u211b", "F": "\u2131", "O": "\U0001d4aa"}

ACCENTS = {"widehat": "\u0302", "hat": "\u0302", "bar": "\u0304",
           "tilde": "\u0303", "widetilde": "\u0303", "vec": "\u20d7",
           "dot": "\u0307"}


def expand_user_macros(s):
    # the replacements contain backslashes, so pass a function, not a template
    for k in sorted(USER_MACROS, key=len, reverse=True):
        v = USER_MACROS[k]
        s = re.sub(re.escape(k) + r"(?![a-zA-Z])", lambda _m, v=v: v, s)
    return s


# ---------------------------------------------------------------------------
# Maths to styled runs
# ---------------------------------------------------------------------------
class R:
    """One styled run of output text."""

    def __init__(self, text, italic=False, bold=False, sub=False, sup=False,
                 mono=False):
        self.text = text
        self.italic = italic
        self.bold = bold
        self.sub = sub
        self.sup = sup
        self.mono = mono

    def restyle(self, **kw):
        return R(self.text, **{**dict(italic=self.italic, bold=self.bold,
                                      sub=self.sub, sup=self.sup,
                                      mono=self.mono), **kw})


def _take_group(s, i):
    """Return (content, next_index) for a {...} group or a single token."""
    while i < len(s) and s[i] in " \t\n":
        i += 1
    if i >= len(s):
        return "", i
    if s[i] == "{":
        depth, j = 1, i + 1
        while j < len(s) and depth:
            if s[j] == "{":
                depth += 1
            elif s[j] == "}":
                depth -= 1
            j += 1
        return s[i + 1:j - 1], j
    if s[i] == "\\":
        m = re.match(r"\\[a-zA-Z]+", s[i:])
        if m:
            return m.group(0), i + m.end()
        return s[i:i + 2], i + 2
    return s[i], i + 1


def math_runs(src, upright=False, bold=False):
    """Convert a LaTeX maths string into styled runs."""
    src = expand_user_macros(src)
    out, i = [], 0
    while i < len(src):
        c = src[i]

        if c == "\\":
            m = re.match(r"\\([a-zA-Z]+)", src[i:])
            if not m:
                if i + 1 < len(src) and src[i + 1] in "{}%&_#$ ,;:!":
                    ch = src[i + 1]
                    out.append(R(" " if ch == " " else
                                 ("" if ch in ",;:!" else ch)))
                    i += 2
                    continue
                i += 1
                continue
            name, i = m.group(1), i + m.end()

            if name in ("mathrm", "mathsf", "text", "textup", "mathtt",
                        "operatorname", "mathbf"):
                g, i = _take_group(src, i)
                out.extend(math_runs(g, upright=True,
                                     bold=(name == "mathbf")))
            elif name == "mathbb":
                g, i = _take_group(src, i)
                out.append(R(BLACKBOARD.get(g.strip(), g.strip())))
            elif name == "mathcal":
                g, i = _take_group(src, i)
                out.append(R(CALLIGRAPHIC.get(g.strip(), g.strip())))
            elif name in ACCENTS:
                g, i = _take_group(src, i)
                inner = math_runs(g, upright=upright, bold=bold)
                if inner:
                    inner[0] = R(inner[0].text + ACCENTS[name],
                                 italic=inner[0].italic, bold=inner[0].bold,
                                 sub=inner[0].sub, sup=inner[0].sup)
                out.extend(inner)
            elif name in ("big", "Big", "bigg", "Bigg", "left", "right",
                          "displaystyle", "textstyle", "limits", "nolimits",
                          "arraybackslash", "hfill", "noexpand"):
                continue
            elif name == "frac":
                a, i = _take_group(src, i)
                b, i = _take_group(src, i)
                out.append(R("("))
                out.extend(math_runs(a, upright, bold))
                out.append(R(")/("))
                out.extend(math_runs(b, upright, bold))
                out.append(R(")"))
            elif name == "sqrt":
                g, i = _take_group(src, i)
                out.append(R("\u221a("))
                out.extend(math_runs(g, upright, bold))
                out.append(R(")"))
            elif name == "pmod":
                g, i = _take_group(src, i)
                out.append(R(" (mod "))
                out.extend(math_runs(g, upright, bold))
                out.append(R(")"))
            elif name in SYMBOLS:
                out.append(R(SYMBOLS[name]))
            else:
                out.append(R(name, italic=False))
            continue

        if c in "_^":
            g, i2 = _take_group(src, i + 1)
            i = i2
            for r in math_runs(g, upright, bold):
                out.append(r.restyle(sub=(c == "_"), sup=(c == "^")))
            continue

        if c == "{":
            g, i = _take_group(src, i)
            out.extend(math_runs(g, upright, bold))
            continue

        if c == "}":
            i += 1
            continue

        if c == "~":
            out.append(R("\u00a0"))
            i += 1
            continue

        if c.isalpha():
            out.append(R(c, italic=not upright, bold=bold))
            i += 1
            continue

        out.append(R(c, bold=bold))
        i += 1
    return out


# ---------------------------------------------------------------------------
# Bibliography
# ---------------------------------------------------------------------------
def parse_bib(path):
    txt = open(path, encoding="utf-8").read()
    entries = {}
    for m in re.finditer(r"@(\w+)\s*\{\s*([^,\s]+)\s*,", txt):
        key = m.group(2)
        start = m.end()
        depth, j = 1, txt.index("{", m.start())
        j += 1
        while j < len(txt) and depth:
            if txt[j] == "{":
                depth += 1
            elif txt[j] == "}":
                depth -= 1
            j += 1
        body = txt[start:j - 1]
        fields = {}
        for fm in re.finditer(r"(\w+)\s*=\s*", body):
            k = fm.group(1).lower()
            p = fm.end()
            if p < len(body) and body[p] == "{":
                d, q = 1, p + 1
                while q < len(body) and d:
                    if body[q] == "{":
                        d += 1
                    elif body[q] == "}":
                        d -= 1
                    q += 1
                fields[k] = body[p + 1:q - 1]
            elif p < len(body) and body[p] == '"':
                q = body.index('"', p + 1)
                fields[k] = body[p + 1:q]
            else:
                q = p
                while q < len(body) and body[q] not in ",\n":
                    q += 1
                fields[k] = body[p:q]
        entries[key] = fields
    return entries


def clean_bib(s):
    s = re.sub(r"[{}]", "", s)
    s = s.replace("\\&", "&").replace("--", "\u2013").replace("~", " ")
    s = re.sub(r"\\[a-zA-Z]+\s*", "", s)
    return re.sub(r"\s+", " ", s).strip()


def format_ref(f):
    au = clean_bib(f.get("author", f.get("editor", "")))
    if au:
        parts = [a.strip() for a in re.split(r"\s+and\s+", au)]
        names = []
        for a in parts:
            if "," in a:
                last, first = a.split(",", 1)
                ini = " ".join(w[0] + "." for w in first.split() if w)
                names.append(f"{last.strip()}, {ini}".strip().rstrip(","))
            else:
                w = a.split()
                if len(w) > 1:
                    names.append(w[-1] + ", " +
                                 " ".join(x[0] + "." for x in w[:-1]))
                else:
                    names.append(a)
        au = ", ".join(names)
    bits = [au, clean_bib(f.get("title", ""))]
    venue = f.get("journal") or f.get("booktitle") or f.get("school") or \
        f.get("publisher") or f.get("howpublished") or ""
    venue = clean_bib(venue)
    if venue:
        bits.append(venue)
    if f.get("volume"):
        v = clean_bib(f["volume"])
        if f.get("number"):
            v += " (" + clean_bib(f["number"]) + ")"
        bits.append(v)
    if f.get("pages"):
        bits.append("pp. " + clean_bib(f["pages"]))
    if f.get("year"):
        bits.append(clean_bib(f["year"]))
    out = ", ".join(b for b in bits if b)
    if not out.endswith("."):
        out += "."
    if f.get("doi"):
        out += " doi:" + clean_bib(f["doi"]) + "."
    return out


# ---------------------------------------------------------------------------
# Document assembly helpers
# ---------------------------------------------------------------------------
def set_base_style(doc, size=12, spacing=2.0):
    st = doc.styles["Normal"]
    st.font.name = BODY_FONT
    st.font.size = Pt(size)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
    pf = st.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.DOUBLE if spacing == 2.0 \
        else WD_LINE_SPACING.ONE_POINT_FIVE
    pf.space_after = Pt(0)


def add_line_numbers(section):
    sectPr = section._sectPr
    ln = OxmlElement("w:lnNumType")
    ln.set(qn("w:countBy"), "1")
    ln.set(qn("w:restart"), "continuous")
    ln.set(qn("w:distance"), "360")
    sectPr.append(ln)


def add_page_numbers(section):
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run()
    for instr in ("begin", "PAGE", "end"):
        el = OxmlElement("w:fldChar" if instr != "PAGE" else "w:instrText")
        if instr == "PAGE":
            el.set(qn("xml:space"), "preserve")
            el.text = " PAGE "
        else:
            el.set(qn("w:fldCharType"), instr)
        r._r.append(el)
    r.font.name = BODY_FONT
    r.font.size = Pt(10)


def emit_runs(par, runs, size=None, base_italic=False):
    for r in runs:
        if not r.text:
            continue
        run = par.add_run(r.text)
        run.italic = r.italic or base_italic
        run.bold = r.bold
        run.font.subscript = r.sub
        run.font.superscript = r.sup
        run.font.name = MONO_FONT if r.mono else BODY_FONT
        if size:
            run.font.size = Pt(size)
    return par


# ---------------------------------------------------------------------------
# Inline (mixed text and maths) rendering
# ---------------------------------------------------------------------------
def inline_runs(text, refs, cites, bold=False, italic=False):
    """Convert body text with $maths$, \\emph{}, \\cite{}, \\ref{} to runs."""
    out = []
    text = text.replace("\\\\", " ")
    i = 0
    while i < len(text):
        c = text[i]

        if c == "$":
            j = text.index("$", i + 1) if "$" in text[i + 1:] else len(text)
            out.extend(math_runs(text[i + 1:j]))
            i = j + 1
            continue

        if c == "\\":
            m = re.match(r"\\([a-zA-Z]+)\*?", text[i:])
            if not m:
                if i + 1 < len(text):
                    ch = text[i + 1]
                    if ch in "%&_#${}":
                        out.append(R(ch, italic=italic, bold=bold))
                        i += 2
                        continue
                    if ch == " ":
                        out.append(R(" "))
                        i += 2
                        continue
                i += 1
                continue
            name = m.group(1)
            i += m.end()

            if name in ("emph", "textit"):
                g, i = _take_group(text, i)
                out.extend(inline_runs(g, refs, cites, bold=bold, italic=True))
            elif name in ("textbf", "textsc"):
                g, i = _take_group(text, i)
                out.extend(inline_runs(g, refs, cites, bold=True,
                                       italic=italic))
            elif name == "texttt":
                g, i = _take_group(text, i)
                out.append(R(re.sub(r"\\", "", g), mono=True))
            elif name.startswith("cite"):
                while i < len(text) and text[i] == "[":
                    _, i = text[i:], i + text[i:].index("]") + 1
                g, i = _take_group(text, i)
                nums = []
                for k in [x.strip() for x in g.split(",") if x.strip()]:
                    if k not in cites:
                        cites.append(k)
                    nums.append(cites.index(k) + 1)
                out.append(R("[" + ",".join(str(n) for n in sorted(nums))
                             + "]"))
            elif name == "ref":
                g, i = _take_group(text, i)
                out.append(R(str(refs.get(g.strip(), "?"))))
            elif name == "label":
                _, i = _take_group(text, i)
            elif name == "S":
                out.append(R("Section "))
            elif name == "url":
                g, i = _take_group(text, i)
                out.append(R(g))
            elif name in ("qed", "noindent", "centering", "raggedright",
                          "arraybackslash", "hfill", "footnotesize",
                          "scriptsize", "small", "normalsize", "toprule",
                          "midrule", "bottomrule", "hline"):
                continue
            elif name == "item":
                continue
            elif name in SYMBOLS:
                out.append(R(SYMBOLS[name]))
            else:
                g, i2 = _take_group(text, i)
                if i2 > i and text[i:i + 1] == "{":
                    i = i2
                    out.extend(inline_runs(g, refs, cites, bold, italic))
            continue

        if c == "~":
            out.append(R("\u00a0", italic=italic, bold=bold))
            i += 1
            continue

        if text.startswith("``", i):
            out.append(R("\u201c", italic=italic, bold=bold))
            i += 2
            continue
        if text.startswith("''", i):
            out.append(R("\u201d", italic=italic, bold=bold))
            i += 2
            continue
        if text.startswith("--", i):
            n = 3 if text.startswith("---", i) else 2
            out.append(R("\u2014" if n == 3 else "\u2013",
                         italic=italic, bold=bold))
            i += n
            continue

        out.append(R(c, italic=italic, bold=bold))
        i += 1

    merged = []
    for r in out:
        if merged and (merged[-1].italic, merged[-1].bold, merged[-1].sub,
                       merged[-1].sup, merged[-1].mono) == \
                (r.italic, r.bold, r.sub, r.sup, r.mono):
            merged[-1].text += r.text
        else:
            merged.append(r)
    return merged


def norm(s):
    return re.sub(r"[ \t]*\n[ \t]*", " ", s).strip()


# ---------------------------------------------------------------------------
# Parse the manuscript into blocks
# ---------------------------------------------------------------------------
def strip_comments(tex):
    out = []
    for ln in tex.splitlines():
        out.append(re.sub(r"(?<!\\)%.*$", "", ln))
    return "\n".join(out)


def grab_env(tex, env, start=0):
    b = "\\begin{%s}" % env
    e = "\\end{%s}" % env
    i = tex.find(b, start)
    if i < 0:
        return None, None, None
    j = tex.find(e, i)
    if j < 0:
        return None, None, None
    return tex[i + len(b):j], i, j + len(e)


ENV_MAP = [("figure*", "figure"), ("figure", "figure"),
           ("table*", "table"), ("table", "table"),
           ("algorithm", "algorithm"), ("itemize", "itemize"),
           ("multline*", "display"), ("gather*", "display"),
           ("align*", "display"),
           ("theorem", "theorem"), ("lemma", "lemma"),
           ("definition", "definition"), ("remark", "remark"),
           ("pf", "proof")]


def parse_blocks(body):
    """Split the body into a flat list of typed blocks."""
    blocks = []
    i, n = 0, len(body)
    while i < n:
        if not body[i:].strip():
            break
        j0 = i + (len(body[i:]) - len(body[i:].lstrip()))
        rest = body[j0:]

        m = re.match(r"\\(section|subsection|subsubsection)(\*?)\s*\{", rest)
        if m:
            g, k = _take_group(body, j0 + rest.index("{"))
            lbl = None
            m2 = re.match(r"\s*\\label\{([^}]*)\}", body[k:])
            if m2:
                lbl = m2.group(1)
                k += m2.end()
            blocks.append({"t": m.group(1), "title": norm(g), "label": lbl,
                           "star": m.group(2) == "*"})
            i = k
            continue

        matched = False
        for env, tag in ENV_MAP:
            if rest.startswith("\\begin{%s}" % env):
                inner, _, end = grab_env(body, env, j0)
                if inner is not None:
                    blocks.append({"t": tag, "env": env, "raw": inner})
                    i = end
                    matched = True
                break
        if matched:
            continue

        if rest.startswith("\\["):
            end = body.index("\\]", j0) + 2
            blocks.append({"t": "display", "env": "[",
                           "raw": body[j0 + 2:end - 2]})
            i = end
            continue

        cands = [body.find("\n\n", j0)]
        for pat in ("\n\\section", "\n\\subsection", "\n\\begin{", "\n\\["):
            cands.append(body.find(pat, j0 + 1))
        cands = [c for c in cands if c > j0]
        j = min(cands) if cands else n
        chunk = body[j0:j].strip()
        if chunk:
            blocks.append({"t": "para", "raw": chunk})
        i = j + 1
    return blocks


FLOAT_COUNTERS = ("table", "figure", "algorithm", "theorem", "lemma",
                  "definition", "remark")


def number_blocks(blocks):
    """Assign numbers to sections, floats and theorem-likes; map labels."""
    refs = {}
    sec = sub = 0
    counters = dict((k, 0) for k in FLOAT_COUNTERS)
    for b in blocks:
        if b["t"] == "section":
            if not b.get("star"):
                sec += 1
                sub = 0
                b["num"] = str(sec)
            else:
                b["num"] = ""
        elif b["t"] == "subsection":
            sub += 1
            b["num"] = "%d.%d" % (sec, sub)
        elif b["t"] in counters:
            counters[b["t"]] += 1
            b["num"] = str(counters[b["t"]])

        if b.get("num"):
            if b.get("label"):
                refs[b["label"]] = b["num"]
            m = re.search(r"\\label\{([^}]*)\}", b.get("raw", ""))
            if m:
                refs[m.group(1)] = b["num"]
    return refs


# ---------------------------------------------------------------------------
# Emit
# ---------------------------------------------------------------------------
def caption_of(raw):
    m = re.search(r"\\caption\{", raw)
    if not m:
        return ""
    g, _ = _take_group(raw, m.end() - 1)
    return norm(g)


def emit_para(doc, text, refs, cites, first=False, size=12, align=None,
              space=6):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space)
    if align is not None:
        p.alignment = align
    else:
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        if not first:
            p.paragraph_format.first_line_indent = Inches(0.25)
    emit_runs(p, inline_runs(norm(text), refs, cites), size=size)
    return p


def emit_display(doc, raw, refs, cites):
    raw = re.sub(r"\\(?:label|notag|nonumber)\{?[^}]*\}?", "", raw)
    for line in re.split(r"\\\\", raw):
        line = line.replace("&", "").strip()
        if not line:
            continue
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
        emit_runs(p, math_runs(norm(line)))


def emit_figure(doc, blk, refs, cites):
    m = re.search(r"\\includegraphics(?:\[[^\]]*\])?\{([^}]*)\}", blk["raw"])
    if m:
        png = os.path.join(ROOT, m.group(1))
        if png.endswith(".pdf"):
            png = png[:-4] + ".png"
        if os.path.exists(png):
            wide = blk.get("env", "").endswith("*")
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(10)
            p.add_run().add_picture(png, width=Inches(6.3 if wide else 4.6))
    cap = caption_of(blk["raw"])
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    p.paragraph_format.space_after = Pt(12)
    r = p.add_run("Fig. %s. " % blk.get("num", "?"))
    r.bold = True
    r.font.size = Pt(10)
    r.font.name = BODY_FONT
    emit_runs(p, inline_runs(cap, refs, cites), size=10)


def parse_tabular(raw):
    m = re.search(r"\\begin\{tabular\*?\}", raw)
    if not m:
        return None
    body = raw[m.end():]
    for _ in range(2):
        if body.lstrip().startswith("{"):
            _g, k = _take_group(body, body.index("{"))
            body = body[k:]
    body = re.split(r"\\end\{tabular", body)[0]
    head, sep, tail = body.partition("\\midrule")
    head = head.replace("\\toprule", "")
    tail = re.split(r"\\bottomrule", tail)[0]
    rows = []
    for part, is_head in ((head, True), (tail, False)):
        for ln in re.split(r"\\\\", part):
            ln = re.sub(r"\\(?:cmidrule|midrule|toprule|bottomrule)"
                        r"(?:\([^)]*\))?(?:\{[^}]*\})?", "", ln).strip()
            if not ln:
                continue
            cells = [c.strip() for c in ln.split("&")]
            rows.append((cells, is_head))
    return rows


def shade(cell, hexcolor):
    el = OxmlElement("w:shd")
    el.set(qn("w:val"), "clear")
    el.set(qn("w:fill"), hexcolor)
    cell._tc.get_or_add_tcPr().append(el)


def emit_table(doc, blk, refs, cites):
    cap = caption_of(blk["raw"])
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    p.paragraph_format.space_before = Pt(10)
    r = p.add_run("Table %s. " % blk.get("num", "?"))
    r.bold = True
    r.font.size = Pt(10)
    r.font.name = BODY_FONT
    emit_runs(p, inline_runs(norm(cap), refs, cites), size=10)

    rows = parse_tabular(blk["raw"])
    if not rows:
        return
    ncol = max(len(c) for c, _ in rows)
    t = doc.add_table(rows=0, cols=ncol)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for cells, is_head in rows:
        row = t.add_row()
        for k in range(ncol):
            cell = row.cells[k]
            cell.paragraphs[0].paragraph_format.space_after = Pt(1)
            cell.paragraphs[0].paragraph_format.line_spacing_rule = \
                WD_LINE_SPACING.SINGLE
            txt = cells[k] if k < len(cells) else ""
            txt = re.sub(r"\\multicolumn\{\d+\}\{[^}]*\}", "", txt)
            runs = inline_runs(norm(txt), refs, cites, bold=is_head)
            emit_runs(cell.paragraphs[0], runs, size=9)
            if is_head:
                shade(cell, "EDEDED")
    doc.add_paragraph().paragraph_format.space_after = Pt(6)


def emit_algorithm(doc, blk, refs, cites):
    cap = caption_of(blk["raw"])
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    p.paragraph_format.space_before = Pt(10)
    r = p.add_run("Algorithm %s. " % blk.get("num", "?"))
    r.bold = True
    r.font.size = Pt(10)
    r.font.name = BODY_FONT
    emit_runs(p, inline_runs(cap, refs, cites), size=10)

    inner, _, _ = grab_env(blk["raw"], "algorithmic")
    if inner is None:
        return
    inner = re.sub(r"^\s*\[\d+\]", "", inner.strip())

    steps, cur = [], ""
    for ln in inner.splitlines():
        if re.match(r"\s*\\(State|Statex|Procedure|EndProcedure|ForAll|"
                    r"EndFor|Return|If|EndIf)", ln):
            if cur.strip():
                steps.append(cur)
            cur = ln
        else:
            cur += " " + ln
    if cur.strip():
        steps.append(cur)

    depth, num = 0, 0
    for st in steps:
        st = st.strip()
        if re.match(r"\\(EndProcedure|EndFor|EndIf)", st):
            depth = max(0, depth - 1)
            continue
        label = ""
        mproc = re.match(r"\\Procedure\{(.*?)\}\{(.*?)\}", st)
        if mproc:
            st = ("procedure " + mproc.group(1) + "(" + mproc.group(2) + ")"
                  + st[mproc.end():])
            indent = depth
            depth += 1
        elif st.startswith("\\ForAll"):
            g, k = _take_group(st, st.index("{"))
            st = "for all " + g + st[k:]
            indent = depth
            depth += 1
        elif st.startswith("\\Statex"):
            st = st[len("\\Statex"):]
            indent = depth
        else:
            st = re.sub(r"^\\State\b", "", st)
            num += 1
            label = "%d:" % num
            indent = depth

        st = re.sub(r"\\Comment\{(.*?)\}", r"   // \1", st, flags=re.S)
        st = re.sub(r"\\Return\b", "return ", st)
        runs = inline_runs(norm(st), refs, cites)
        if not any(x.text.strip() for x in runs) and not label:
            continue
        p = doc.add_paragraph()
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.left_indent = Inches(0.30 + 0.22 * indent)
        if label:
            rr = p.add_run(label + " ")
            rr.font.name = MONO_FONT
            rr.font.size = Pt(9)
        emit_runs(p, runs, size=9.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(6)


def emit_theoremlike(doc, blk, refs, cites):
    kind = {"theorem": "Theorem", "lemma": "Lemma",
            "definition": "Definition", "remark": "Remark",
            "proof": "Proof"}[blk["t"]]
    raw = blk["raw"]
    title = ""
    if re.match(r"\s*\[", raw):
        j = raw.index("[")
        d, k = 1, j + 1
        while k < len(raw) and d:
            if raw[k] == "[":
                d += 1
            elif raw[k] == "]":
                d -= 1
            k += 1
        title = raw[j + 1:k - 1]
        raw = raw[k:]
    raw = re.sub(r"\\label\{[^}]*\}", "", raw)

    head = kind if kind == "Proof" else "%s %s" % (kind, blk.get("num", ""))
    if title:
        head += " (" + title.replace("$", "") + ")"
    head += "."

    chunks = [c for c in re.split(r"\n\s*\n", raw) if c.strip()]
    if not chunks:
        chunks = [""]
    for idx, ch in enumerate(chunks):
        ch = ch.strip()
        if ch.startswith("\\[") or ch.startswith("\\begin{multline"):
            if ch.startswith("\\["):
                inner = ch[2:ch.rindex("\\]")] if "\\]" in ch else ch[2:]
            else:
                got, _, _ = grab_env(ch, "multline*")
                inner = got if got is not None else ch
            emit_display(doc, inner, refs, cites)
            continue
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.space_after = Pt(6)
        if idx == 0:
            r = p.add_run(head + " ")
            r.bold = True
            r.font.name = BODY_FONT
        emit_runs(p, inline_runs(norm(ch), refs, cites),
                  base_italic=blk["t"] in ("theorem", "lemma"))


def emit_itemize(doc, raw, refs, cites):
    for it in [x for x in re.split(r"\\item\b", raw) if x.strip()]:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(6)
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        emit_runs(p, inline_runs(norm(it), refs, cites))


# ---------------------------------------------------------------------------
# Front matter and driver
# ---------------------------------------------------------------------------
AUTHORS = [
    ("Md Mazharul Islam", "0000-0001-6225-1487",
     "mazharul.islam1@northsouth.edu", True),
    ("Mubasshir Ahmed", "0000-0002-7502-3835",
     "mubasshir.ahmed@northsouth.edu", False),
    ("Md Redowan Zaman Anik", "", "redowan.anik@northsouth.edu", False),
]
AFFIL = ("Department of Electrical and Computer Engineering, "
         "North South University, Dhaka, Bangladesh")

CREDIT = [
    ("Md Mazharul Islam", "Conceptualization, Methodology, Software, "
     "Formal analysis, Investigation, Validation, Writing \u2013 original "
     "draft, Writing \u2013 review and editing"),
    ("Mubasshir Ahmed", "Software, Data curation, Investigation, "
     "Writing \u2013 review and editing"),
    ("Md Redowan Zaman Anik", "Software, Data curation, Visualization, "
     "Validation, Writing \u2013 review and editing"),
]


def heading(doc, text, size=13, before=14, after=6, bold=True):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    r = p.add_run(text)
    r.bold = bold
    r.font.size = Pt(size)
    r.font.name = BODY_FONT
    return p


def build(argv=None):
    tex = strip_comments(open(TEX, encoding="utf-8").read())
    bib = parse_bib(BIB)

    title = norm(_take_group(tex, tex.index("{", tex.index(
        "\\title[mode=title]")))[0])
    abstract, _, _ = grab_env(tex, "abstract")
    highlights, _, _ = grab_env(tex, "highlights")
    keywords, _, _ = grab_env(tex, "keywords")

    body = tex.split("\\maketitle", 1)[1]
    body = body.split("\\bibliographystyle", 1)[0]
    body = body.replace("\\printcredits", "")

    blocks = parse_blocks(body)
    refs = number_blocks(blocks)
    cites = []

    doc = Document()
    set_base_style(doc)
    sec = doc.sections[0]
    sec.top_margin = sec.bottom_margin = Inches(1.0)
    sec.left_margin = sec.right_margin = Inches(1.0)
    add_line_numbers(sec)
    add_page_numbers(sec)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(10)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    r = p.add_run(title)
    r.bold = True
    r.font.size = Pt(15)
    r.font.name = BODY_FONT

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    for k, (name, orcid, mail, corr) in enumerate(AUTHORS):
        if k:
            p.add_run(", ").font.size = Pt(12)
        rr = p.add_run(name + ("*" if corr else ""))
        rr.font.size = Pt(12)
        rr.font.name = BODY_FONT

    for line, sz in ((AFFIL, 11),
                     ("* Corresponding author: %s" % AUTHORS[0][2], 10)):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
        rr = p.add_run(line)
        rr.font.size = Pt(sz)
        rr.font.name = BODY_FONT
        rr.italic = (sz == 11)

    heading(doc, "Abstract")
    emit_para(doc, abstract, refs, cites, first=True)

    heading(doc, "Highlights")
    emit_itemize(doc, highlights, refs, cites)

    heading(doc, "Keywords")
    kw = "; ".join(x.strip() for x in keywords.split("\\sep") if x.strip())
    emit_para(doc, kw, refs, cites, first=True)

    doc.add_paragraph().add_run().add_break()

    prev_head = False
    for blk in blocks:
        t = blk["t"]
        if t == "section":
            num = blk.get("num")
            heading(doc, (num + ". " if num else "") + blk["title"], size=13)
            prev_head = True
        elif t == "subsection":
            heading(doc, blk["num"] + ". " + blk["title"], size=12,
                    before=10)
            prev_head = True
        elif t == "subsubsection":
            heading(doc, blk["title"], size=12, before=8, bold=False)
            prev_head = True
        elif t == "para":
            emit_para(doc, blk["raw"], refs, cites, first=prev_head)
            prev_head = False
        elif t == "display":
            emit_display(doc, blk["raw"], refs, cites)
            prev_head = True
        elif t == "figure":
            emit_figure(doc, blk, refs, cites)
            prev_head = True
        elif t == "table":
            emit_table(doc, blk, refs, cites)
            prev_head = True
        elif t == "algorithm":
            emit_algorithm(doc, blk, refs, cites)
            prev_head = True
        elif t == "itemize":
            emit_itemize(doc, blk["raw"], refs, cites)
            prev_head = True
        elif t in ("theorem", "lemma", "definition", "remark", "proof"):
            emit_theoremlike(doc, blk, refs, cites)
            prev_head = True

    heading(doc, "CRediT authorship contribution statement")
    for name, roles in CREDIT:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.space_after = Pt(6)
        r = p.add_run(name + ": ")
        r.bold = True
        r.font.name = BODY_FONT
        r2 = p.add_run(roles + ".")
        r2.font.name = BODY_FONT

    heading(doc, "References")
    for k, key in enumerate(cites, 1):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
        p.paragraph_format.left_indent = Inches(0.35)
        p.paragraph_format.first_line_indent = Inches(-0.35)
        r = p.add_run("[%d] " % k)
        r.font.size = Pt(10.5)
        r.font.name = BODY_FONT
        r2 = p.add_run(format_ref(bib.get(key, {"title": key})))
        r2.font.size = Pt(10.5)
        r2.font.name = BODY_FONT

    doc.save(OUT)
    print("wrote", OUT)
    print("  sections numbered:",
          sum(1 for b in blocks if b["t"] == "section" and b.get("num")))
    print("  tables:", sum(1 for b in blocks if b["t"] == "table"),
          " figures:", sum(1 for b in blocks if b["t"] == "figure"),
          " algorithms:", sum(1 for b in blocks if b["t"] == "algorithm"))
    print("  references cited:", len(cites))


if __name__ == "__main__":
    build()
