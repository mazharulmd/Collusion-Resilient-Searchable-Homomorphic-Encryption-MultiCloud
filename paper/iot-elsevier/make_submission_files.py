"""Build the separate submission files for Elsevier Internet of Things:
highlights (.txt/.docx) and cover letter (.docx). Highlights are read from
the manuscript so the two never disagree; each is checked against the
85-character limit."""
import re

from docx import Document
from docx.shared import Pt

TEX = "CRSHE_IoT.tex"
src = open(TEX).read()
block = src[src.index(r"\begin{highlights}"):src.index(r"\end{highlights}")]
items = []
for raw in block.split(r"\item")[1:]:
    t = " ".join(raw.split())
    t = re.sub(r"\$t=n-1\$", "t = n-1", t)
    items.append(t)
assert 3 <= len(items) <= 5
for t in items:
    assert len(t) <= 85, (len(t), t)

with open("CRSHE_IoT_Highlights.txt", "w") as f:
    f.write("Highlights\n\n" + "\n".join(f"- {t}" for t in items) + "\n")
d = Document()
d.add_heading("Highlights", level=1)
for t in items:
    d.add_paragraph(t, style="List Bullet")
d.save("CRSHE_IoT_Highlights.docx")

TITLE = ("CR-SHE: Oblivious homomorphic aggregation for collusion-resilient search "
         "and analytics over encrypted multi-cloud IoT data")
letter = [
    "Md Mazharul Islam",
    "Department of Electrical and Computer Engineering, North South University, Dhaka, Bangladesh",
    "mazharul.islam1@northsouth.edu",
    "[DATE]",
    "",
    "The Editors",
    "Internet of Things (Elsevier)",
    "",
    "Dear Editors,",
    "",
    f"We submit our manuscript entitled “{TITLE}” for consideration as a "
    "full-length research article in Internet of Things.",
    "",
    "IoT deployments increasingly replicate telemetry across independent cloud providers "
    "for availability, and the data must remain encrypted once it leaves the device domain. "
    "Existing distributed-trust systems hide which records a query selects, but either "
    "compute over plaintext, stop at retrieval, or require the servers to interact. CR-SHE "
    "makes both keyword retrieval and analytics over the matched records oblivious to a "
    "coalition of up to n-1 colluding providers, over data encrypted under a single "
    "owner's key, with no communication between providers.",
    "",
    "The manuscript contributes (i) a construction in which each provider folds its share "
    "of the selection vector into encrypted columns and returns one ciphertext, so the "
    "matched set and its size never leave the client; (ii) a proof that, under these "
    "constraints, a homomorphic layer is necessary, since an additively shared store "
    "cannot answer such queries locally; (iii) a fast path for registered aggregates, a "
    "two-round flexible path for conjunctions and ad-hoc weights, and an index mask that "
    "removes posting-list leakage at rest; and (iv) security proofs against n-1 collusion "
    "and an evaluation on two public IoT corpora, including a direct comparison with a "
    "single server holding the same encrypted columns.",
    "",
    "We believe the paper suits Internet of Things because its central trade-off is an IoT "
    "one: measured against a single-server design, the multi-cloud trust split saves a "
    "modest factor in server work but more than three orders of magnitude in uplink and "
    "client cost, which is what thin, metered IoT clients need.",
    "",
    "This manuscript is original, has not been published, and is not under consideration "
    "elsewhere. A version of the manuscript and the code are publicly available in the "
    "authors' GitHub repository. All authors have approved the submission and declare no "
    "competing interests.",
    "",
    "Yours sincerely,",
    "Md Mazharul Islam, on behalf of all authors",
]
d = Document()
for line in letter:
    p = d.add_paragraph(line)
    for r in p.runs:
        r.font.size = Pt(11)
d.save("CRSHE_IoT_CoverLetter.docx")
print("highlights:", [len(t) for t in items])


# ---------------------------------------------------------------------------
# Title page and declaration of competing interests (Word, as Editorial
# Manager asks). Bracketed, highlighted text is for the authors to complete;
# nothing there is filled in on their behalf.
# ---------------------------------------------------------------------------
from docx.enum.text import WD_COLOR_INDEX

AUTHORS = [
    ("Md Mazharul Islam", "mazharul.islam1@northsouth.edu", "0000-0001-6225-1487", True),
    ("Mubasshir Ahmed", "mubasshir.ahmed@northsouth.edu", "0000-0002-7502-3835", False),
    ("Md Redowan Zaman Anik", "redowan.anik@northsouth.edu", None, False),
]
AFFIL = ("Department of Electrical and Computer Engineering, North South University, "
         "Dhaka, Bangladesh")

# CRediT roles are taken from the manuscript source, so the two cannot drift.
credits = re.findall(r"\\author\[1\]\{([^}]*)\}.*?\\credit\{(.*?)\}", src, re.S)
credits = [(n, " ".join(c.replace("--", "\u2013").split())) for n, c in credits]
assert len(credits) == 3


def todo(par, text):
    run = par.add_run(text)
    run.font.highlight_color = WD_COLOR_INDEX.YELLOW
    return run


d = Document()
d.add_heading(TITLE, level=1)
p = d.add_paragraph()
for k, (name, _, _, corr) in enumerate(AUTHORS):
    p.add_run(name)
    sup = p.add_run("a" + (",*" if corr else ""))
    sup.font.superscript = True
    if k < len(AUTHORS) - 1:
        p.add_run(", ")
p = d.add_paragraph()
p.add_run("a").font.superscript = True
p.add_run(AFFIL)
p = d.add_paragraph()
p.add_run("*").font.superscript = True
p.add_run(f"Corresponding author: {AUTHORS[0][0]}, {AUTHORS[0][1]}")

d.add_heading("Author e-mail addresses and ORCID iDs", level=2)
for name, mail, orcid, _ in AUTHORS:
    d.add_paragraph(f"{name}: {mail}" + (f"; ORCID {orcid}" if orcid else ""), style="List Bullet")

d.add_heading("Acknowledgements", level=2)
d.add_paragraph("This research did not receive any specific grant from funding agencies in the "
                "public, commercial, or not-for-profit sectors.")

d.add_heading("CRediT authorship contribution statement", level=2)
for name, roles in credits:
    p = d.add_paragraph()
    p.add_run(name + ": ").bold = True
    p.add_run(roles + ".")

d.save("CRSHE_IoT_TitlePage.docx")

d = Document()
d.add_heading("Declaration of interests", level=1)
d.add_paragraph(f"Manuscript title: {TITLE}")
d.add_paragraph("\u2612 The authors declare that they have no known competing financial interests "
                "or personal relationships that could have appeared to influence the work "
                "reported in this paper.")
p = d.add_paragraph()
todo(p, "[Confirm before uploading:] ")
p.add_run("No author currently serves, or has previously served, as an editor or editorial "
          "board member of Internet of Things.")
d.add_paragraph("On behalf of all authors: " + AUTHORS[0][0])
d.save("CRSHE_IoT_DeclarationOfInterest.docx")
print("title page and declaration written")
